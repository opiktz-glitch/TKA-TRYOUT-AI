from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent
OUTPUT = ROOT / "Panduan_Deploy_Server_LAN.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color="D9D9D9"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), color)


def set_run_font(run, size=None, bold=None, color=None):
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def paragraph(doc, text="", style=None, before=0, after=6):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.12
    if text:
        set_run_font(p.add_run(text), 10.5)
    return p


def code_block(doc, lines):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.18)
    p_pr = p._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F3F4F6")
    p_pr.append(shading)
    for index, line in enumerate(lines):
        if index:
            p.add_run("\n")
        run = p.add_run(line)
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
        run._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
        run.font.size = Pt(9)


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    set_run_font(p.add_run(text), 10.5)


def numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(3)
    set_run_font(p.add_run(text), 10.5)


def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(14 if level == 1 else 9)
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run(text)
    set_run_font(run, 14 if level == 1 else 11.5, bold=True, color=(0, 0, 0))
    return p


def add_footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Panduan Deploy Server LAN - TKA TryOut AI")
    set_run_font(run, 8, color=(90, 90, 90))


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    add_footer(section)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(6)
    set_run_font(title.add_run("Panduan Deploy Server LAN"), 22, bold=True, color=(0, 0, 0))
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(18)
    set_run_font(subtitle.add_run("TKA TryOut AI dengan Docker"), 11, color=(80, 80, 80))

    paragraph(
        doc,
        "Dokumen ini menjelaskan cara menjalankan TKA TryOut AI pada satu komputer server di jaringan sekolah dan mengaksesnya dari HP atau laptop lain pada Wi-Fi atau LAN yang sama. Ikuti urutan konfigurasi, migrasi, dan verifikasi sebelum digunakan untuk tryout.",
        after=10,
    )

    heading(doc, "1. Prasyarat")
    for item in (
        "Docker Desktop sudah berjalan di komputer server.",
        "Komputer server dan perangkat pengguna berada pada jaringan LAN atau Wi-Fi yang sama.",
        "Windows Firewall mengizinkan koneksi TCP masuk ke port 5173 dan 8000 untuk jaringan privat.",
        "Folder proyek TKA TryOut AI tersedia pada komputer server.",
    ):
        bullet(doc, item)
    paragraph(doc, "Periksa alamat IP server dengan PowerShell:", after=3)
    code_block(doc, ["ipconfig"])
    paragraph(doc, "Gunakan nilai IPv4 Address pada adapter Wi-Fi atau Ethernet yang aktif. Contoh dalam panduan ini menggunakan 172.16.64.216. Ganti dengan IP server Anda bila berbeda atau berubah.")

    heading(doc, "2. Konfigurasi Environment")
    heading(doc, "Root .env untuk frontend", 2)
    paragraph(doc, "Dari root proyek, salin template konfigurasi Docker:", after=3)
    code_block(doc, ["Copy-Item .env.example .env"])
    paragraph(doc, "Isi file .env dengan alamat API yang dapat diakses browser pengguna:", after=3)
    code_block(doc, ["VITE_API_URL=http://172.16.64.216:8000"])
    paragraph(doc, "Jangan gunakan localhost untuk server LAN. Pada HP atau laptop pengguna, localhost berarti perangkat pengguna itu sendiri, bukan komputer server.")

    heading(doc, "backend .env untuk API", 2)
    paragraph(doc, "Jika belum ada, salin template backend:", after=3)
    code_block(doc, ["Copy-Item backend/.env.example backend/.env"])
    paragraph(doc, "Buat SECRET_KEY acak yang kuat:", after=3)
    code_block(doc, ['python -c "import secrets; print(secrets.token_urlsafe(64))"'])
    paragraph(doc, "Tempel hasilnya pada SECRET_KEY di backend/.env. Jangan membagikan atau memasukkan file .env ke Git. Akses dari alamat IP privat LAN telah didukung oleh konfigurasi aplikasi. Untuk domain publik di masa depan, isi CORS_ORIGINS dengan domain frontend.")

    heading(doc, "3. Migrasi Database")
    paragraph(doc, "Jika database sudah pernah digunakan sebelum perbaikan race condition attempt, jalankan script ini sekali sebelum aplikasi menerima pengguna:", after=3)
    code_block(doc, ["cd backend", "python scripts/create_attempt_index.py", "cd .."])
    paragraph(doc, "Script menambahkan index unik agar satu siswa tidak memiliki dua attempt IN_PROGRESS untuk tryout yang sama.")

    heading(doc, "4. Menjalankan Aplikasi")
    paragraph(doc, "Jalankan perintah berikut dari root proyek:", after=3)
    code_block(doc, ["docker compose up -d --build", "docker compose ps"])
    paragraph(doc, "Gunakan --build setiap kali nilai VITE_API_URL di root .env berubah. Nilai tersebut dibundel ke frontend saat proses build; docker compose restart saja tidak memperbaruinya.")

    heading(doc, "5. Alamat Akses")
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.autofit = False
    table.columns[0].width = Inches(1.45)
    table.columns[1].width = Inches(4.9)
    header = table.rows[0].cells
    for cell, label in zip(header, ("Layanan", "Alamat")):
        set_cell_shading(cell, "1F4E78")
        set_cell_border(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(p.add_run(label), 10, bold=True, color=(255, 255, 255))
    for service, address in (
        ("Frontend", "http://172.16.64.216:5173"),
        ("API status", "http://172.16.64.216:8000/api/system/status"),
    ):
        cells = table.add_row().cells
        for index, value in enumerate((service, address)):
            set_cell_border(cells[index])
            cells[index].vertical_alignment = 1
            set_run_font(cells[index].paragraphs[0].add_run(value), 10)
    paragraph(doc, "Buka frontend dari perangkat lain dalam jaringan yang sama. Dokumentasi API dinonaktifkan pada mode Docker atau production.", before=6)

    heading(doc, "6. Verifikasi Setelah Deploy")
    for item in (
        "Buka halaman frontend dari HP atau laptop selain komputer server, lalu login.",
        "Mulai tryout, simpan jawaban, dan submit tryout.",
        "Pastikan hasil dan riwayat tryout ditampilkan dengan benar.",
        "Jalankan backup manual dari Pengaturan Admin bila tersedia.",
        "Periksa log jika terjadi masalah.",
    ):
        numbered(doc, item)
    code_block(doc, ["docker compose logs --tail=100"])

    heading(doc, "7. Backup dan Operasi Harian")
    for item in (
        "Database utama berjalan di Docker named volume.",
        "Hasil backup SQLite disimpan pada backend/backups di komputer server.",
        "Salin folder backup secara berkala ke media atau cloud terpisah.",
        "Uji pemulihan backup pada lingkungan salinan, bukan pada database aktif.",
    ):
        bullet(doc, item)
    code_block(doc, [
        "# Melihat status container",
        "docker compose ps",
        "",
        "# Melihat log berjalan",
        "docker compose logs -f",
        "",
        "# Menghentikan aplikasi tanpa menghapus data volume",
        "docker compose down",
        "",
        "# Membangun ulang setelah pembaruan kode atau konfigurasi",
        "docker compose up -d --build",
    ])

    heading(doc, "8. Catatan Keamanan")
    for item in (
        "Jangan membuka port 5173 dan 8000 langsung ke internet publik untuk setup LAN ini.",
        "Jangan commit file .env, database SQLite, atau folder backup.",
        "Gunakan password admin yang kuat dan simpan SECRET_KEY dengan aman.",
        "Untuk akses dari internet, gunakan domain HTTPS dan reverse proxy. Konfigurasi LAN ini bukan pengganti setup publik.",
    ):
        bullet(doc, item)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
